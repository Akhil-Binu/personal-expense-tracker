"""
Excel service for Personal Expense Tracker.
Handles generating formatted multi-sheet Excel reports and importing expense records.
"""
import io
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database

# Brand Colors for Excel
HEADER_BG = "1E293B"       # Slate 800
HEADER_TEXT = "FFFFFF"     # White
ACCENT_BG = "0D9488"       # Teal 600
ACCENT_LIGHT = "F0FDFA"    # Teal 50
ZEBRA_BG = "F8FAFC"        # Slate 50
TOTAL_BG = "E2E8F0"        # Slate 200
BORDER_COLOR = "CBD5E1"    # Slate 300

THIN_BORDER = Border(
    left=Side(style='thin', color=BORDER_COLOR),
    right=Side(style='thin', color=BORDER_COLOR),
    top=Side(style='thin', color=BORDER_COLOR),
    bottom=Side(style='thin', color=BORDER_COLOR)
)

TOTAL_BORDER = Border(
    top=Side(style='thin', color="475569"),
    bottom=Side(style='double', color="1E293B")
)

def export_expenses_to_excel(month_str=None, category_id=None, payment_method=None, search=None, currency_symbol=None, user_id=1):
    """
    Generates a beautifully formatted multi-sheet Excel workbook (.xlsx).
    Returns a io.BytesIO buffer.
    """
    user_id = user_id or 1
    if not currency_symbol:
        currency_symbol = database.get_currency()["symbol"]
    curr_fmt = f'"{currency_symbol}"#,##0.00'

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    # ----------------------------------------------------
    # Sheet 1: Summary Dashboard
    # ----------------------------------------------------
    ws_summary = wb.create_sheet(title="Executive Summary")
    ws_summary.views.sheetView[0].showGridLines = True
    
    metrics = database.get_monthly_metrics(month_str, user_id=user_id)
    category_breakdown = database.get_category_breakdown(month_str, user_id=user_id)
    payment_methods = database.get_payment_method_breakdown(month_str, user_id=user_id)
    
    # Title Banner
    ws_summary.merge_cells("A1:G1")
    title_cell = ws_summary["A1"]
    title_cell.value = f"Personal Expense Report — {metrics['month']}"
    title_cell.font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 40
    
    # Subtitle
    ws_summary.merge_cells("A2:G2")
    sub_cell = ws_summary["A2"]
    sub_cell.value = f"Generated on {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
    sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[2].height = 20
    
    # KPI Cards (Row 4-5)
    kpis = [
        ("Total Spent", f"{currency_symbol}{metrics['total_spend']:,.2f}", "0D9488"),
        ("Monthly Budget", f"{currency_symbol}{metrics['total_budget']:,.2f}", "3B82F6"),
        ("Remaining Budget", f"{currency_symbol}{(metrics['total_budget'] - metrics['total_spend']):,.2f}", "10B981" if metrics['total_budget'] >= metrics['total_spend'] else "EF4444"),
        ("Daily Average", f"{currency_symbol}{metrics['daily_average']:,.2f}", "8B5CF6"),
        ("Total Transactions", str(metrics['expense_count']), "F59E0B"),
        ("Top Category", metrics['top_category'], "EC4899"),
    ]
    
    ws_summary.row_dimensions[4].height = 20
    ws_summary.row_dimensions[5].height = 28
    
    col_idx = 1
    for title, val, color in kpis:
        col_letter = get_column_letter(col_idx)
        c_title = ws_summary[f"{col_letter}4"]
        c_title.value = title
        c_title.font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
        c_title.alignment = Alignment(horizontal="center", vertical="center")
        c_title.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        c_title.border = THIN_BORDER
        
        c_val = ws_summary[f"{col_letter}5"]
        c_val.value = val
        c_val.font = Font(name="Segoe UI", size=13, bold=True, color=color)
        c_val.alignment = Alignment(horizontal="center", vertical="center")
        c_val.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        c_val.border = THIN_BORDER
        col_idx += 1
        
    # Category Breakdown Header
    ws_summary.cell(row=7, column=1, value="Category Spending Breakdown").font = Font(name="Segoe UI", size=12, bold=True, color=HEADER_BG)
    
    cat_headers = ["Category", "Monthly Budget", "Actual Spent", "Remaining", "Budget Used %", "Share of Total %", "Tx Count"]
    ws_summary.row_dimensions[8].height = 26
    for c_i, h in enumerate(cat_headers, start=1):
        cell = ws_summary.cell(row=8, column=c_i, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=ACCENT_BG, end_color=ACCENT_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_i > 1 else "left", vertical="center")
        cell.border = THIN_BORDER

    start_cat_row = 9
    current_row = start_cat_row
    for cat in category_breakdown:
        ws_summary.row_dimensions[current_row].height = 22
        is_zebra = (current_row % 2 == 0)
        row_bg = ZEBRA_BG if is_zebra else "FFFFFF"
        fill_style = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
        
        ws_summary.cell(row=current_row, column=1, value=cat["name"]).alignment = Alignment(horizontal="left", vertical="center")
        
        c_budget = ws_summary.cell(row=current_row, column=2, value=cat["monthly_budget"])
        c_budget.number_format = curr_fmt
        c_budget.alignment = Alignment(horizontal="right", vertical="center")
        
        c_spent = ws_summary.cell(row=current_row, column=3, value=cat["spent"])
        c_spent.number_format = curr_fmt
        c_spent.alignment = Alignment(horizontal="right", vertical="center")
        
        # Remaining formula
        c_rem = ws_summary.cell(row=current_row, column=4, value=f"=B{current_row}-C{current_row}")
        c_rem.number_format = curr_fmt
        c_rem.alignment = Alignment(horizontal="right", vertical="center")
        
        # Budget used formula
        c_used = ws_summary.cell(row=current_row, column=5, value=f"=IF(B{current_row}>0, C{current_row}/B{current_row}, 0)")
        c_used.number_format = "0.0%"
        c_used.alignment = Alignment(horizontal="right", vertical="center")
        
        c_share = ws_summary.cell(row=current_row, column=6, value=f"={cat['percentage_of_total'] / 100:.4f}")
        c_share.number_format = "0.0%"
        c_share.alignment = Alignment(horizontal="right", vertical="center")
        
        c_cnt = ws_summary.cell(row=current_row, column=7, value=cat["transaction_count"])
        c_cnt.alignment = Alignment(horizontal="center", vertical="center")
        
        for c_i in range(1, 8):
            cell = ws_summary.cell(row=current_row, column=c_i)
            cell.font = Font(name="Segoe UI", size=10)
            cell.fill = fill_style
            cell.border = THIN_BORDER
            
        current_row += 1
        
    # Category Totals Row
    total_cat_row = current_row
    ws_summary.row_dimensions[total_cat_row].height = 24
    ws_summary.cell(row=total_cat_row, column=1, value="Total").alignment = Alignment(horizontal="left", vertical="center")
    
    c_tot_b = ws_summary.cell(row=total_cat_row, column=2, value=f"=SUM(B{start_cat_row}:B{total_cat_row-1})")
    c_tot_b.number_format = curr_fmt
    
    c_tot_s = ws_summary.cell(row=total_cat_row, column=3, value=f"=SUM(C{start_cat_row}:C{total_cat_row-1})")
    c_tot_s.number_format = curr_fmt
    
    c_tot_r = ws_summary.cell(row=total_cat_row, column=4, value=f"=SUM(D{start_cat_row}:D{total_cat_row-1})")
    c_tot_r.number_format = curr_fmt
    
    c_tot_u = ws_summary.cell(row=total_cat_row, column=5, value=f"=IF(B{total_cat_row}>0, C{total_cat_row}/B{total_cat_row}, 0)")
    c_tot_u.number_format = "0.0%"
    
    c_tot_sh = ws_summary.cell(row=total_cat_row, column=6, value=f"=SUM(F{start_cat_row}:F{total_cat_row-1})")
    c_tot_sh.number_format = "0.0%"
    
    c_tot_cnt = ws_summary.cell(row=total_cat_row, column=7, value=f"=SUM(G{start_cat_row}:G{total_cat_row-1})")
    
    for c_i in range(1, 8):
        cell = ws_summary.cell(row=total_cat_row, column=c_i)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
        cell.fill = PatternFill(start_color=TOTAL_BG, end_color=TOTAL_BG, fill_type="solid")
        cell.border = TOTAL_BORDER
        if c_i in [2, 3, 4, 5, 6]:
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif c_i == 7:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
    # Payment Method Breakdown Section
    pm_start_row = total_cat_row + 3
    ws_summary.cell(row=pm_start_row - 1, column=1, value="Payment Methods Overview").font = Font(name="Segoe UI", size=12, bold=True, color=HEADER_BG)
    pm_headers = ["Payment Method", f"Total Spent ({currency_symbol})", "Transactions", "Share of Total"]
    for i, h in enumerate(pm_headers, start=1):
        cell = ws_summary.cell(row=pm_start_row, column=i, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if i > 1 else "left", vertical="center")
        cell.border = THIN_BORDER
        
    pm_curr_row = pm_start_row + 1
    total_pm_spent = sum(p["total"] for p in payment_methods) if payment_methods else 1.0
    for p in payment_methods:
        ws_summary.cell(row=pm_curr_row, column=1, value=p["payment_method"]).border = THIN_BORDER
        c_amt = ws_summary.cell(row=pm_curr_row, column=2, value=p["total"])
        c_amt.number_format = curr_fmt
        c_amt.alignment = Alignment(horizontal="right", vertical="center")
        c_amt.border = THIN_BORDER
        
        c_ct = ws_summary.cell(row=pm_curr_row, column=3, value=p["count"])
        c_ct.alignment = Alignment(horizontal="center", vertical="center")
        c_ct.border = THIN_BORDER
        
        share_val = (p["total"] / total_pm_spent) if total_pm_spent > 0 else 0
        c_sh = ws_summary.cell(row=pm_curr_row, column=4, value=round(share_val, 4))
        c_sh.number_format = "0.0%"
        c_sh.alignment = Alignment(horizontal="right", vertical="center")
        c_sh.border = THIN_BORDER
        
        pm_curr_row += 1

    # Adjust Summary Columns Width
    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 16)
    ws_summary.column_dimensions['A'].width = 24

    # ----------------------------------------------------
    # Sheet 2: Detailed Expense Records
    # ----------------------------------------------------
    ws_details = wb.create_sheet(title="All Transactions")
    ws_details.views.sheetView[0].showGridLines = True
    
    # Header Banner
    headers = ["ID", "Date", "Category", "Description", "Payment Method", f"Amount ({currency_symbol})", "Notes"]
    ws_details.row_dimensions[1].height = 28
    for c_idx, h in enumerate(headers, start=1):
        cell = ws_details.cell(row=1, column=c_idx, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 2, 5] else ("right" if c_idx == 6 else "left"), vertical="center")
        cell.border = THIN_BORDER

    # Fetch expenses with filters applied
    start_date = f"{month_str}-01" if month_str else None
    end_date = f"{month_str}-31" if month_str else None
    expenses = database.get_expenses(start_date=start_date, end_date=end_date, category_id=category_id, payment_method=payment_method, search=search, user_id=user_id)
    
    row_num = 2
    for exp in expenses:
        ws_details.row_dimensions[row_num].height = 20
        is_zebra = (row_num % 2 == 0)
        row_bg = ZEBRA_BG if is_zebra else "FFFFFF"
        fill_style = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
        
        c_id = ws_details.cell(row=row_num, column=1, value=exp["id"])
        c_id.alignment = Alignment(horizontal="center", vertical="center")
        
        c_date = ws_details.cell(row=row_num, column=2, value=exp["date"])
        c_date.alignment = Alignment(horizontal="center", vertical="center")
        
        c_cat = ws_details.cell(row=row_num, column=3, value=exp["category_name"])
        c_cat.alignment = Alignment(horizontal="left", vertical="center")
        
        c_desc = ws_details.cell(row=row_num, column=4, value=exp["description"])
        c_desc.alignment = Alignment(horizontal="left", vertical="center")
        
        c_pm = ws_details.cell(row=row_num, column=5, value=exp["payment_method"])
        c_pm.alignment = Alignment(horizontal="center", vertical="center")
        
        c_amt = ws_details.cell(row=row_num, column=6, value=exp["amount"])
        c_amt.number_format = curr_fmt
        c_amt.alignment = Alignment(horizontal="right", vertical="center")
        
        c_note = ws_details.cell(row=row_num, column=7, value=exp.get("notes") or "")
        c_note.alignment = Alignment(horizontal="left", vertical="center")
        
        for c_idx in range(1, 8):
            cell = ws_details.cell(row=row_num, column=c_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.fill = fill_style
            cell.border = THIN_BORDER
            
        row_num += 1

    # Total Row for Transactions
    if row_num > 2:
        ws_details.row_dimensions[row_num].height = 24
        ws_details.cell(row=row_num, column=5, value="Total Spending").alignment = Alignment(horizontal="right", vertical="center")
        c_tot = ws_details.cell(row=row_num, column=6, value=f"=SUM(F2:F{row_num-1})")
        c_tot.number_format = curr_fmt
        c_tot.alignment = Alignment(horizontal="right", vertical="center")
        
        for c_idx in range(1, 8):
            cell = ws_details.cell(row=row_num, column=c_idx)
            cell.font = Font(name="Segoe UI", size=10, bold=True)
            cell.fill = PatternFill(start_color=TOTAL_BG, end_color=TOTAL_BG, fill_type="solid")
            cell.border = TOTAL_BORDER
            
    # Enable AutoFilter on Transactions
    ws_details.auto_filter.ref = f"A1:G{max(2, row_num - 1)}"
    
    # Adjust Detailed Columns Width
    col_widths = {
        'A': 8,   # ID
        'B': 13,  # Date
        'C': 20,  # Category
        'D': 30,  # Description
        'E': 20,  # Payment Method
        'F': 16,  # Amount
        'G': 28   # Notes
    }
    for col_letter, width in col_widths.items():
        ws_details.column_dimensions[col_letter].width = width

    # Save to memory buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_import_template():
    """Generates an Excel template (.xlsx) with sample rows to guide users when importing."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expense Import Template"
    ws.views.sheetView[0].showGridLines = True
    
    headers = ["Date (YYYY-MM-DD)", "Category", "Amount", "Payment Method", "Description", "Notes"]
    ws.row_dimensions[1].height = 26
    for c_i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_i, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
        
    sample_rows = [
        ("2026-10-24", "Food & Dining", 18.50, "Credit Card", "Quick lunch", "Deli sandwich"),
        ("2026-10-24", "Groceries", 45.20, "Debit Card", "Supermarket vegetables", "Fresh produce"),
        ("2026-10-23", "Transportation", 25.00, "UPI / Digital Wallet", "Fuel recharge", "Car gas"),
        ("2026-10-22", "Entertainment", 14.99, "Credit Card", "Movie ticket", "Cinema night"),
    ]
    
    for r_idx, row in enumerate(sample_rows, start=2):
        ws.row_dimensions[r_idx].height = 20
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = THIN_BORDER
            if c_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 3:
                cell.number_format = "$#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx == 4:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 24
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 25
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def import_expenses_from_excel(file_stream, user_id=1):
    """
    Parses an uploaded Excel file and inserts valid expense rows into SQLite for a user.
    Returns: dict with imported count, skipped count, and list of errors if any.
    """
    user_id = user_id or 1
    wb = openpyxl.load_workbook(file_stream, data_only=True)
    ws = wb.active
    
    categories = {c["name"].strip().lower(): c["id"] for c in database.get_categories(user_id=user_id)}
    
    imported_count = 0
    skipped_count = 0
    errors = []
    
    # Detect header positions
    header_map = {}
    for col in range(1, ws.max_column + 1):
        val = str(ws.cell(row=1, column=col).value or "").strip().lower()
        if "date" in val:
            header_map["date"] = col
        elif "cat" in val:
            header_map["category"] = col
        elif "amount" in val or "price" in val or "cost" in val:
            header_map["amount"] = col
        elif "payment" in val or "method" in val:
            header_map["payment"] = col
        elif "desc" in val or "item" in val or "title" in val:
            header_map["description"] = col
        elif "note" in val or "comment" in val:
            header_map["notes"] = col

    # Fallback to column indices if headers weren't named identically
    date_col = header_map.get("date", 1)
    cat_col = header_map.get("category", 2)
    amt_col = header_map.get("amount", 3)
    pay_col = header_map.get("payment", 4)
    desc_col = header_map.get("description", 5)
    notes_col = header_map.get("notes", 6)
    
    for row_idx in range(2, ws.max_row + 1):
        raw_date = ws.cell(row=row_idx, column=date_col).value
        raw_cat = ws.cell(row=row_idx, column=cat_col).value
        raw_amt = ws.cell(row=row_idx, column=amt_col).value
        raw_pay = ws.cell(row=row_idx, column=pay_col).value
        raw_desc = ws.cell(row=row_idx, column=desc_col).value
        raw_notes = ws.cell(row=row_idx, column=notes_col).value
        
        # Skip completely blank rows
        if not any([raw_date, raw_cat, raw_amt, raw_desc]):
            continue
            
        try:
            # Parse Date
            if isinstance(raw_date, (datetime, date)):
                exp_date = raw_date.strftime("%Y-%m-%d")
            else:
                date_str = str(raw_date).strip()
                # Try common formats
                parsed_dt = None
                for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d"):
                    try:
                        parsed_dt = datetime.strptime(date_str, fmt)
                        break
                    except ValueError:
                        pass
                if parsed_dt:
                    exp_date = parsed_dt.strftime("%Y-%m-%d")
                else:
                    exp_date = date.today().isoformat()
                    
            # Parse Amount
            amt_clean = str(raw_amt).replace("$", "").replace(",", "").strip()
            amount = float(amt_clean)
            if amount <= 0:
                raise ValueError("Amount must be greater than 0")
                
            # Parse Category
            cat_name = str(raw_cat or "Others").strip()
            cat_id = categories.get(cat_name.lower())
            if not cat_id:
                # Create category if missing
                cat_id = database.add_category(name=cat_name, icon="tag", color="#6B7280", monthly_budget=100.0, user_id=user_id)
                categories[cat_name.lower()] = cat_id
                
            payment_method = str(raw_pay or "Cash").strip()
            description = str(raw_desc or "Imported Expense").strip()
            notes = str(raw_notes or "").strip()
            
            database.add_expense(
                expense_date=exp_date,
                category_id=cat_id,
                amount=amount,
                payment_method=payment_method,
                description=description,
                notes=notes,
                user_id=user_id
            )
            imported_count += 1
        except Exception as e:
            skipped_count += 1
            errors.append(f"Row {row_idx}: {str(e)}")
            
    return {
        "imported": imported_count,
        "skipped": skipped_count,
        "errors": errors[:5]  # return first 5 error samples
    }
